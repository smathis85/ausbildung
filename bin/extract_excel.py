#!/usr/bin/env python3
"""Read-only source extraction. Run with bundled Python/openpyxl; output is PRIVATE."""
import sys, json, hashlib, datetime, re, collections, os
import openpyxl

def text(v):
    return str(v).strip() if v is not None else ''

def date(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime('%Y-%m-%d')
    raw = text(v)
    match = re.search(r'(\d{1,2})\D{1,3}(\d{1,2})\.(\d{4})', raw)
    if match:
        d, m, y = map(int, match.groups())
        return datetime.date(y, m, d).isoformat()
    return None

def extract(path):
    wb = openpyxl.load_workbook(path, data_only=False)
    people, sessions, enrollments, attendance, warnings = {}, [], [], [], []
    previous = {}
    summary = []
    for sheet in wb:
        if not sheet.title.isdigit():
            continue
        year = int(sheet.title)
        headers = {text(sheet.cell(4,c).value):c for c in range(1,sheet.max_column+1) if sheet.cell(4,c).value is not None}
        carry_col, total_col = headers['Bisherige Termine'], headers['Teilnahmen gesamt']
        first_col = carry_col + 1
        numbered = collections.Counter()
        columns = {}
        for c in range(first_col, total_col):
            value = sheet.cell(4,c).value
            if value is None:
                continue
            day = date(value)
            label = day or text(value)
            numbered[label] += 1
            key = f'{year}:{c}'
            columns[c] = key
            sessions.append(dict(source_key=key, year=year, date=day, unit=numbered[label],
                title='Ausbildung' if day and isinstance(value,datetime.datetime) else text(value),
                source_label=text(value), source_cell=f'{year}!{sheet.cell(4,c).coordinate}'))
        count = marks = 0
        seen = set()
        for r in range(5,sheet.max_row+1):
            last, first, department = [text(sheet.cell(r,c).value) for c in (1,2,3)]
            if not (last and first and department):
                continue
            key = hashlib.sha256(('\x1f'.join(x.casefold() for x in (last,first,department))).encode()).hexdigest()
            if key in seen:
                raise ValueError(f'Doppelte Person im selben Jahr: {year}, Zeile {r}; manuell klären')
            seen.add(key)
            people[key] = dict(source_key=key,last_name=last,first_name=first,department=department)
            count += 1
            carry = sheet.cell(r,carry_col).value
            if carry is None: carry=0
            if not isinstance(carry,(int,float)) or int(carry)!=carry or carry<0:
                raise ValueError(f'Ungültiger Vorjahresstand: {year}, Zeile {r}')
            n = 0
            for c in range(first_col,total_col):
                v = sheet.cell(r,c).value
                if v is None: continue
                if text(v).lower() != 'x' or c not in columns:
                    raise ValueError(f'Unklarer Anwesenheitseintrag: {year}!{sheet.cell(r,c).coordinate}')
                attendance.append(dict(person=key,session=columns[c]))
                n += 1
            marks += n
            prior = previous.get(key,0)
            delta = int(carry)-prior
            if key in previous and delta != 0:
                warnings.append(dict(type='carry_adjustment',person=key,year=year,source_row=r,difference=delta))
            begin=sheet.cell(r,4).value
            course=text(sheet.cell(r,5).value) if carry_col==6 else ''
            exam=sheet.cell(r,total_col+1).value
            reported=sheet.cell(r,total_col+2).value
            comment=text(sheet.cell(r,total_col+3).value)
            enrollments.append(dict(person=key,year=year,adjustment=delta,source_carry=int(carry),source_total=int(carry)+n,
                start_label=text(begin),start_date=date(begin),course=course,exam_date=date(exam),exam_label=text(exam),
                reported=text(reported),comment=comment,source_row=r,
                adjustment_reason='Excel-Vortrag vor erstem erfassten Jahr' if key not in previous else 'Abgleich mit dem Vorjahresstand der Excel-Liste'))
            previous[key]=int(carry)+n
        summary.append(dict(year=year,people=count,attendance=marks))
    return dict(format_version=1,source_sha256=hashlib.sha256(open(path,'rb').read()).hexdigest(),
        people=list(people.values()),sessions=sessions,enrollments=enrollments,attendance=attendance,warnings=warnings,summary=summary)

if __name__=='__main__':
    result=extract(sys.argv[1])
    fd=os.open(sys.argv[2],os.O_WRONLY|os.O_CREAT|os.O_TRUNC,0o600)
    with os.fdopen(fd,'w') as f: json.dump(result,f,ensure_ascii=False)
    print(json.dumps(dict(people=len(result['people']),sessions=len(result['sessions']),carry_adjustments=len(result['warnings']),years=result['summary']),ensure_ascii=False))
